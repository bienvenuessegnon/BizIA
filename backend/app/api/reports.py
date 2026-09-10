from datetime import UTC, datetime
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from reportlab.graphics.shapes import Circle, Drawing, Line, PolyLine, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.api.auth import current_user
from app.services.store import get_user_store
from app.utils.errors import ApiError

router = APIRouter()

_BLUE = colors.HexColor("#2563EB")
_NAVY = colors.HexColor("#0F172A")
_SLATE = colors.HexColor("#475569")
_LIGHT = colors.HexColor("#F1F5F9")
_GREEN = colors.HexColor("#0F766E")
_RED = colors.HexColor("#B91C1C")
_AMBER = colors.HexColor("#D97706")


def _report_payload(analysis: dict) -> dict:
    return {
        "kpis": analysis.get("kpis"),
        "insights": analysis.get("insights"),
        "alerts": analysis.get("alerts"),
        "recommendations": analysis.get("recommendations"),
        "top_sold": analysis.get("top_sold"),
        "top_profit": analysis.get("top_profit"),
        "low_stock": analysis.get("low_stock"),
        "trend": analysis.get("trend"),
        "anomalies": analysis.get("anomalies"),
        "week_over_week": analysis.get("week_over_week"),
    }


def _money(value: object) -> str:
    try:
        amount = float(value or 0)
    except (TypeError, ValueError):
        amount = 0
    return f"{amount:,.0f}".replace(",", "\u00a0") + "\u00a0FCFA"


def _number(value: object) -> str:
    try:
        return f"{float(value or 0):,.0f}".replace(",", "\u00a0")
    except (TypeError, ValueError):
        return "0"


def _text(value: object) -> str:
    """Texte utilisateur/LLM sûr pour le mini-HTML de ReportLab."""
    return escape(str(value or "")).replace("\n", "<br/>")


def _lines(report: dict) -> list[str]:
    kpis = report.get("kpis") or {}
    lines = [
        f"Chiffre d'affaires : {kpis.get('revenue', 0):,.0f} FCFA",
        f"Coûts : {kpis.get('cost', 0):,.0f} FCFA",
        f"Bénéfice : {kpis.get('profit', 0):,.0f} FCFA",
        f"Marge : {kpis.get('margin_pct', 0):.1f} %",
        "",
        "Insights",
        *[f"• {item}" for item in report.get("insights") or []],
        "",
        "Recommandations",
        *[
            f"• [{item.get('priority', '')}] {item.get('action', '')} — {item.get('why', '')}"
            for item in report.get("recommendations") or []
        ],
    ]
    return lines


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "BizTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=28,
            textColor=_NAVY,
            spaceAfter=4 * mm,
        ),
        "subtitle": ParagraphStyle(
            "BizSubtitle",
            parent=base["BodyText"],
            fontSize=10,
            leading=14,
            textColor=_SLATE,
            spaceAfter=7 * mm,
        ),
        "h1": ParagraphStyle(
            "BizH1",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=_NAVY,
            spaceBefore=5 * mm,
            spaceAfter=3 * mm,
        ),
        "h2": ParagraphStyle(
            "BizH2",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=_NAVY,
            spaceBefore=2 * mm,
            spaceAfter=1.5 * mm,
        ),
        "body": ParagraphStyle(
            "BizBody",
            parent=base["BodyText"],
            fontSize=9.5,
            leading=14,
            textColor=_SLATE,
            spaceAfter=2 * mm,
        ),
        "small": ParagraphStyle(
            "BizSmall",
            parent=base["BodyText"],
            fontSize=8,
            leading=10,
            textColor=_SLATE,
        ),
        "kpi_label": ParagraphStyle(
            "BizKpiLabel",
            parent=base["BodyText"],
            fontSize=7.5,
            leading=9,
            textColor=_SLATE,
        ),
        "kpi_value": ParagraphStyle(
            "BizKpiValue",
            parent=base["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=16,
            textColor=_NAVY,
        ),
    }


def _footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#CBD5E1"))
    canvas.line(18 * mm, 14 * mm, A4[0] - 18 * mm, 14 * mm)
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(_SLATE)
    canvas.drawString(18 * mm, 9 * mm, "BizIA · Rapport confidentiel")
    canvas.drawRightString(A4[0] - 18 * mm, 9 * mm, f"Page {document.page}")
    canvas.restoreState()


def _kpi_table(report: dict, styles: dict[str, ParagraphStyle]) -> Table:
    kpis = report.get("kpis") or {}
    values = [
        ("Chiffre d'affaires", _money(kpis.get("revenue"))),
        ("Coûts", _money(kpis.get("cost"))),
        ("Bénéfice", _money(kpis.get("profit"))),
        ("Marge", f"{float(kpis.get('margin_pct') or 0):.1f}\u00a0%"),
        ("Unités vendues", _number(kpis.get("units_sold"))),
        ("Nombre de ventes", _number(kpis.get("sales_count"))),
    ]
    cells = [
        [
            [
                Paragraph(_text(label), styles["kpi_label"]),
                Spacer(1, 1.5 * mm),
                Paragraph(_text(value), styles["kpi_value"]),
            ]
            for label, value in values[index : index + 3]
        ]
        for index in (0, 3)
    ]
    table = Table(cells, colWidths=[55 * mm] * 3, rowHeights=[21 * mm] * 2)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), _LIGHT),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.white),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3 * mm),
            ]
        )
    )
    return table


def _executive_summary(report: dict) -> str:
    kpis = report.get("kpis") or {}
    profit = float(kpis.get("profit") or 0)
    margin = float(kpis.get("margin_pct") or 0)
    direction = "positif" if profit >= 0 else "déficitaire"
    top = (report.get("top_profit") or [{}])[0]
    top_text = (
        f" Le produit le plus rentable est <b>{_text(top.get('name'))}</b>, "
        f"avec {_money(top.get('profit'))} de bénéfice."
        if top
        else ""
    )
    return (
        f"L'activité présente un résultat <b>{direction}</b> de <b>{_money(profit)}</b>, "
        f"soit une marge de <b>{margin:.1f}\u00a0%</b>.{top_text}"
    )


def _trend_chart(points: list[dict]) -> Drawing:
    width, height = 480, 170
    drawing = Drawing(width, height)
    left, bottom, plot_width, plot_height = 52, 30, 410, 120
    drawing.add(Line(left, bottom, left, bottom + plot_height, strokeColor=colors.HexColor("#CBD5E1")))
    drawing.add(Line(left, bottom, left + plot_width, bottom, strokeColor=colors.HexColor("#CBD5E1")))

    values = [float(point.get("revenue") or 0) for point in points]
    maximum = max(values or [1]) or 1
    for index in range(4):
        y = bottom + plot_height * index / 3
        value = maximum * index / 3
        drawing.add(Line(left, y, left + plot_width, y, strokeColor=colors.HexColor("#E2E8F0")))
        drawing.add(String(left - 5, y - 3, _number(value), fontSize=6.5, textAnchor="end", fillColor=_SLATE))

    denominator = max(len(points) - 1, 1)
    coordinates = [
        (left + plot_width * index / denominator, bottom + plot_height * value / maximum)
        for index, value in enumerate(values)
    ]
    if len(coordinates) == 1:
        coordinates.append((left + plot_width, coordinates[0][1]))
    drawing.add(PolyLine(coordinates, strokeColor=_BLUE, strokeWidth=2.2))
    label_step = max(1, len(points) // 5)
    for index, (x, y) in enumerate(coordinates[: len(points)]):
        drawing.add(Circle(x, y, 2.5, fillColor=_BLUE, strokeColor=colors.white))
        if index % label_step == 0 or index == len(points) - 1:
            label = str(points[index].get("period") or "")[-10:]
            drawing.add(String(x, 14, label, fontSize=6.5, textAnchor="middle", fillColor=_SLATE))
    return drawing


def _ranking_chart(items: list[dict]) -> Drawing:
    items = items[:5]
    width, height = 480, max(90, 27 * len(items) + 20)
    drawing = Drawing(width, height)
    maximum = max([float(item.get("revenue") or 0) for item in items] or [1]) or 1
    for index, item in enumerate(items):
        y = height - 25 - index * 27
        name = str(item.get("name") or item.get("sku") or "Produit")
        if len(name) > 25:
            name = name[:22] + "…"
        drawing.add(String(0, y + 3, name, fontSize=7.5, fillColor=_NAVY))
        bar_width = 250 * float(item.get("revenue") or 0) / maximum
        drawing.add(Rect(135, y, max(bar_width, 1), 11, fillColor=_GREEN, strokeColor=None))
        drawing.add(String(395, y + 2, _money(item.get("revenue")), fontSize=7, fillColor=_SLATE))
    return drawing


def _data_table(
    headers: list[str],
    rows: list[list[object]],
    widths: list[float],
    styles: dict[str, ParagraphStyle],
) -> Table:
    data = [[Paragraph(f"<b>{_text(value)}</b>", styles["small"]) for value in headers]]
    data.extend([[Paragraph(_text(value), styles["small"]) for value in row] for row in rows])
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 2 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2 * mm),
            ]
        )
    )
    return table


def _bullet(text: object, styles: dict[str, ParagraphStyle], color: colors.Color = _BLUE) -> Table:
    marker = Paragraph("●", ParagraphStyle("BulletMarker", parent=styles["body"], textColor=color))
    content = Paragraph(_text(text), styles["body"])
    return Table([[marker, content]], colWidths=[5 * mm, 160 * mm], style=[("VALIGN", (0, 0), (-1, -1), "TOP")])


def _pdf(report: dict) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        title="Rapport d'analyse business BizIA",
        author="BizIA",
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=17 * mm,
        bottomMargin=20 * mm,
    )
    styles = _styles()
    generated = datetime.now(UTC).strftime("%d/%m/%Y à %H:%M UTC")
    story = [
        Paragraph("Rapport d'analyse business", styles["title"]),
        Paragraph(f"BizIA · Généré le {generated} · Données de la dernière analyse validée", styles["subtitle"]),
        Paragraph("Bilan exécutif", styles["h1"]),
        Paragraph(_executive_summary(report), styles["body"]),
        Spacer(1, 2 * mm),
        _kpi_table(report, styles),
    ]

    trend = report.get("trend") or []
    if trend:
        story.extend(
            [
                Paragraph("Évolution du chiffre d'affaires", styles["h1"]),
                Paragraph("Revenus observés sur la période analysée.", styles["body"]),
                _trend_chart(trend),
            ]
        )

    top_sold = report.get("top_sold") or []
    if top_sold:
        story.extend(
            [
                Paragraph("Produits moteurs du chiffre d'affaires", styles["h1"]),
                _ranking_chart(top_sold),
            ]
        )

    top_profit = report.get("top_profit") or []
    if top_profit:
        story.extend(
            [
                Paragraph("Classement de rentabilité", styles["h1"]),
                _data_table(
                    ["Produit", "Unités", "Chiffre d'affaires", "Bénéfice"],
                    [
                        [
                            item.get("name") or item.get("sku"),
                            _number(item.get("units_sold")),
                            _money(item.get("revenue")),
                            _money(item.get("profit")),
                        ]
                        for item in top_profit[:10]
                    ],
                    [55 * mm, 24 * mm, 43 * mm, 43 * mm],
                    styles,
                ),
            ]
        )

    story.append(PageBreak())
    story.append(Paragraph("Analyse, conseils et vigilance", styles["title"]))
    insights = report.get("insights") or []
    story.append(Paragraph("Ce qu'il faut retenir", styles["h1"]))
    if insights:
        story.extend(_bullet(item, styles) for item in insights)
    else:
        story.append(Paragraph("Aucun insight supplémentaire pour cette analyse.", styles["body"]))

    recommendations = report.get("recommendations") or []
    story.append(Paragraph("Conseils et actions recommandées", styles["h1"]))
    if recommendations:
        for item in recommendations:
            priority = str(item.get("priority") or "").upper()
            color = _RED if priority == "HIGH" else (_AMBER if priority == "MEDIUM" else _GREEN)
            story.append(
                KeepTogether(
                    [
                        Paragraph(
                            f"<b>{_text(item.get('action'))}</b> · Priorité {_text(priority or 'NORMALE')}",
                            styles["h2"],
                        ),
                        _bullet(item.get("why"), styles, color),
                    ]
                )
            )
    else:
        story.append(Paragraph("Aucune action prioritaire détectée.", styles["body"]))

    alerts = report.get("alerts") or []
    anomalies = report.get("anomalies") or []
    if alerts or anomalies:
        story.append(Paragraph("Alertes et anomalies", styles["h1"]))
        rows = [
            [
                str(item.get("severity") or "").upper(),
                item.get("title"),
                item.get("detail"),
            ]
            for item in alerts
        ]
        rows.extend(
            [
                str(item.get("severity") or "").upper(),
                f"Anomalie · {item.get('period') or ''}",
                item.get("message"),
            ]
            for item in anomalies
        )
        story.append(
            _data_table(
                ["Niveau", "Sujet", "Détail"],
                rows,
                [22 * mm, 48 * mm, 95 * mm],
                styles,
            )
        )

    low_stock = report.get("low_stock") or []
    if low_stock:
        story.extend(
            [
                Paragraph("Stocks à surveiller", styles["h1"]),
                _data_table(
                    ["Produit", "Stock restant", "Seuil d'alerte"],
                    [
                        [
                            item.get("name") or item.get("sku"),
                            _number(item.get("stock_quantity")),
                            _number(item.get("low_stock_threshold")),
                        ]
                        for item in low_stock
                    ],
                    [85 * mm, 40 * mm, 40 * mm],
                    styles,
                ),
            ]
        )

    story.extend(
        [
            Spacer(1, 7 * mm),
            Paragraph(
                "Note : les indicateurs financiers sont calculés à partir des données validées. "
                "Les recommandations constituent une aide à la décision et doivent être adaptées "
                "au contexte réel de l'entreprise.",
                styles["small"],
            ),
        ]
    )
    document.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return output.getvalue()


def _docx(report: dict) -> bytes:
    output = BytesIO()
    document = Document()
    document.add_heading("Rapport d'analyse BizIA", level=0)
    for line in _lines(report):
        if line in {"Insights", "Recommandations"}:
            document.add_heading(line, level=1)
        elif line:
            document.add_paragraph(line)
    document.save(output)
    return output.getvalue()


@router.post("/generate")
def generate_report(
    format: str = Query(default="pdf", pattern="^(pdf|docx)$"),
    user: dict = Depends(current_user),
):
    analysis = get_user_store(user["id"]).get_last_analysis()
    if analysis is None:
        raise ApiError(
            400,
            "no_analysis",
            "Lancez une analyse avant d'exporter un rapport.",
        )
    report = _report_payload(analysis)
    if format == "pdf":
        return Response(
            _pdf(report),
            media_type="application/pdf",
            headers={"Content-Disposition": 'attachment; filename="bizia-rapport.pdf"'},
        )
    if format == "docx":
        return Response(
            _docx(report),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="bizia-rapport.docx"'},
        )
    raise ApiError(422, "unsupported_report_format", "Choisissez PDF ou Word.")
